# -*- coding: utf-8 -*-
"""Slides da faculdade: "faz um slide sobre X com 8 slides e manda pra mim no zap".

1. O roteiro: a IA (Gemini) monta título, slides (com o layout de cada um), notas do apresentador e referências,
   usando o que a internet diz agora sobre o tema (sem chave, pelo ferramentas.web_snippets).
2. O arquivo, de dois jeitos:
   - padrão, grátis e sem conta: um .pptx caprichado montado aqui mesmo (python-pptx), com capa, divisórias,
     tópicos, duas colunas, destaque e citação, e as cores do estilo escolhido;
   - com o Google conectado ("conecta o Google"): direto no Google Slides, e o PDF vem exportado pelo Drive.
3. Salva em Documentos\\Slides, abre para o senhor ver e, se pedir, manda pelo WhatsApp (pergunta antes).
"""

import json
import re
import unicodedata
from datetime import datetime
from pathlib import Path

LAYOUTS = ("topicos", "duas_colunas", "destaque", "citacao", "secao")
STYLES = {
    # fundo, fundo do card, destaque, texto, texto suave
    "escuro": {"bg": "0B1630", "card": "13234A", "accent": "35C8FF", "accent2": "7B61FF", "text": "FFFFFF", "muted": "A9BCDF"},
    "claro": {"bg": "F7F9FC", "card": "FFFFFF", "accent": "1F6FEB", "accent2": "7B61FF", "text": "142033", "muted": "5B6B85"},
    "vermelho": {"bg": "1A0508", "card": "2B0B10", "accent": "FF5A52", "accent2": "FFB347", "text": "FFFFFF", "muted": "E3B3B3"},
    "dourado": {"bg": "141006", "card": "241C0A", "accent": "FFC94D", "accent2": "FF8A3D", "text": "FFFFFF", "muted": "DCCB9E"},
}

OUTLINE_PROMPT = (
    "Você monta apresentações de faculdade, em português do Brasil, bonitas e com conteúdo de verdade (nada raso). "
    "Responda SÓ um JSON neste formato:\n"
    '{"titulo":"título forte","subtitulo":"uma linha (matéria, contexto)",'
    '"slides":[{"layout":"topicos|duas_colunas|destaque|citacao|secao","titulo":"título curto",'
    '"topicos":["frase curta e informativa"],"colunas":[{"titulo":"...","topicos":["..."]},{"titulo":"...","topicos":["..."]}],'
    '"destaque":"um número, data ou frase de impacto","legenda":"o que o destaque significa",'
    '"notas":"o que o apresentador fala neste slide (3 a 5 frases)"}],'
    '"referencias":["autor/site, título, ano"]}\n'
    "Regras: use exatamente o número de slides pedido (sem contar a capa e as referências); comece com uma visão "
    "geral, varie os layouts (a maioria topicos, com 1 ou 2 de duas_colunas, destaque e citacao quando couber; "
    "secao só para dividir partes grandes); 3 a 5 tópicos por slide, cada um com no máximo 18 palavras e com "
    "informação concreta (datas, nomes, números, causas, consequências); termine com um slide de conclusão; "
    "use as fontes recebidas nas referências e não invente dados."
)


def _say(text):
    print(f"Jarvis > {text}")


def norm(text):
    text = unicodedata.normalize("NFKD", str(text or "").lower())
    return "".join(c for c in text if not unicodedata.combining(c))


def safe_name(text):
    name = re.sub(r'[<>:"/\\|?*\n\r\t]', " ", str(text or "")).strip(" .")
    return re.sub(r"\s+", " ", name)[:80] or f"Slides {datetime.now():%d-%m-%Y %H-%M}"


# ========================= ROTEIRO =========================

def clean_outline(data, topic, count):
    """Confere e arruma o roteiro que a IA mandou (layout desconhecido, tópico vazio, número de slides)."""
    if not isinstance(data, dict):
        raise ValueError("a IA não devolveu o roteiro")
    slides = []
    for s in data.get("slides") or []:
        if not isinstance(s, dict):
            continue
        layout = s.get("layout") if s.get("layout") in LAYOUTS else "topicos"
        item = {"layout": layout, "titulo": str(s.get("titulo") or "").strip()[:90],
                "topicos": [str(t).strip()[:220] for t in (s.get("topicos") or []) if str(t).strip()][:6],
                "notas": str(s.get("notas") or "").strip()[:1500],
                "destaque": str(s.get("destaque") or "").strip()[:120], "legenda": str(s.get("legenda") or "").strip()[:200]}
        cols = [c for c in (s.get("colunas") or []) if isinstance(c, dict)][:2]
        item["colunas"] = [{"titulo": str(c.get("titulo") or "").strip()[:60],
                            "topicos": [str(t).strip()[:160] for t in (c.get("topicos") or []) if str(t).strip()][:5]}
                           for c in cols]
        if layout == "duas_colunas" and len(item["colunas"]) < 2:
            item["layout"] = "topicos"
        if layout in ("destaque", "citacao") and not item["destaque"]:
            item["layout"] = "topicos"
        if item["layout"] == "topicos" and not item["topicos"]:
            continue
        if item["titulo"] or item["destaque"]:
            slides.append(item)
    if not slides:
        raise ValueError("o roteiro veio vazio")
    if count:
        slides = slides[:int(count)]
    return {"titulo": str(data.get("titulo") or topic).strip()[:120] or topic,
            "subtitulo": str(data.get("subtitulo") or "").strip()[:160],
            "slides": slides,
            "referencias": [str(r).strip()[:200] for r in (data.get("referencias") or []) if str(r).strip()][:8]}


def outline(topic, count, level, think, snippets=None, extra=""):
    """O roteiro da apresentação, montado pela IA com as fontes de agora."""
    facts = ""
    try:
        found = snippets(f"{topic}", 8) if snippets else []
        facts = "\n".join(f"- {r['titulo']} ({r['site']}): {r['trecho']}" for r in found)
    except Exception:
        facts = ""
    ask = (f"Tema: {topic}\nNúmero de slides: {count}\nNível: {level or 'faculdade'}\n"
           + (f"Pedido do aluno: {extra}\n" if extra else "")
           + (f"Fontes de agora (use nas referências):\n{facts}\n" if facts else ""))
    raw = think(OUTLINE_PROMPT, ask, json_mode=True, temperature=0.4)
    data = raw if isinstance(raw, dict) else json.loads(re.search(r"\{.*\}", str(raw), re.S).group(0))
    return clean_outline(data, topic, count)


# ========================= .PPTX CAPRICHADO =========================

def make_pptx(plan, path, style="escuro"):
    """Monta o .pptx (16:9) a partir do roteiro. Devolve o caminho."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.util import Inches, Pt, Emu

    c = STYLES.get(norm(style), STYLES["escuro"])
    rgb = lambda h: RGBColor.from_string(h)  # noqa: E731
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]
    FONT = "Segoe UI"

    def background(slide):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(c["bg"])

    def box(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, line=None):
        sh = slide.shapes.add_shape(shape, x, y, w, h)
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(color)
        if line:
            sh.line.color.rgb = rgb(line)
            sh.line.width = Pt(1)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def text(slide, x, y, w, h, value, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = str(value)
        run.font.size, run.font.bold, run.font.name = Pt(size), bold, FONT
        run.font.color.rgb = rgb(color)
        return tf

    def bullets(slide, x, y, w, h, items, size=20):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(18)
            dot = p.add_run()
            dot.text = "●  "
            dot.font.size, dot.font.name = Pt(size - 6), FONT
            dot.font.color.rgb = rgb(c["accent"])
            run = p.add_run()
            run.text = item
            run.font.size, run.font.name = Pt(size), FONT
            run.font.color.rgb = rgb(c["text"])
        return tf

    def header(slide, title, n):
        box(slide, Inches(0.6), Inches(0.55), Inches(0.12), Inches(0.75), c["accent"])
        text(slide, Inches(0.9), Inches(0.45), W - Inches(2.2), Inches(0.95), title, 32, c["text"], bold=True,
             anchor=MSO_ANCHOR.MIDDLE)
        text(slide, W - Inches(1.4), H - Inches(0.6), Inches(0.9), Inches(0.4), f"{n:02d}", 12, c["muted"],
             align=PP_ALIGN.RIGHT)
        box(slide, Inches(0.6), H - Inches(0.42), W - Inches(1.2), Emu(12000), c["card"])

    def notes(slide, value):
        if value:
            slide.notes_slide.notes_text_frame.text = value

    # capa
    s = prs.slides.add_slide(blank)
    background(s)
    box(s, 0, 0, Inches(0.35), H, c["accent"])
    box(s, W - Inches(4.2), Inches(-1.2), Inches(5.5), Inches(5.5), c["card"], MSO_SHAPE.OVAL)
    box(s, W - Inches(2.6), Inches(3.9), Inches(3.4), Inches(3.4), c["accent2"], MSO_SHAPE.OVAL)
    text(s, Inches(1.0), Inches(2.3), Inches(9.5), Inches(2.2), plan["titulo"], 48, c["text"], bold=True,
         anchor=MSO_ANCHOR.BOTTOM)
    box(s, Inches(1.05), Inches(4.65), Inches(1.6), Emu(60000), c["accent"])
    if plan.get("subtitulo"):
        text(s, Inches(1.0), Inches(4.85), Inches(9.5), Inches(1.0), plan["subtitulo"], 20, c["muted"])
    text(s, Inches(1.0), H - Inches(1.0), Inches(6), Inches(0.5), datetime.now().strftime("%d/%m/%Y"), 14, c["muted"])

    for n, sl in enumerate(plan["slides"], 1):
        s = prs.slides.add_slide(blank)
        background(s)
        lay = sl["layout"]
        if lay == "secao":
            box(s, 0, 0, W, H, c["card"])
            text(s, Inches(1.0), Inches(2.4), Inches(11), Inches(1.6), sl["titulo"], 44, c["text"], bold=True,
                 anchor=MSO_ANCHOR.MIDDLE)
            box(s, Inches(1.05), Inches(4.15), Inches(2.2), Emu(60000), c["accent"])
            if sl["topicos"]:
                text(s, Inches(1.0), Inches(4.4), Inches(11), Inches(1.2), sl["topicos"][0], 20, c["muted"])
        elif lay == "destaque":
            header(s, sl["titulo"], n)
            text(s, Inches(0.9), Inches(2.0), W - Inches(1.8), Inches(2.2), sl["destaque"], 72, c["accent"], bold=True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            if sl["legenda"]:
                text(s, Inches(1.5), Inches(4.3), W - Inches(3.0), Inches(1.0), sl["legenda"], 22, c["text"],
                     align=PP_ALIGN.CENTER)
            if sl["topicos"]:
                text(s, Inches(1.5), Inches(5.3), W - Inches(3.0), Inches(1.2), " · ".join(sl["topicos"][:3]), 16,
                     c["muted"], align=PP_ALIGN.CENTER)
        elif lay == "citacao":
            header(s, sl["titulo"], n)
            text(s, Inches(1.0), Inches(1.6), Inches(1.2), Inches(1.4), "“", 120, c["accent"], bold=True)
            text(s, Inches(1.9), Inches(2.2), W - Inches(3.6), Inches(2.6), sl["destaque"], 30, c["text"],
                 anchor=MSO_ANCHOR.MIDDLE)
            if sl["legenda"]:
                text(s, Inches(1.9), Inches(5.0), W - Inches(3.6), Inches(0.7), f"— {sl['legenda']}", 18, c["muted"])
        elif lay == "duas_colunas":
            header(s, sl["titulo"], n)
            colw = (W - Inches(1.2) - Inches(0.4)) / 2
            for i, col in enumerate(sl["colunas"]):
                x = Inches(0.6) + i * (colw + Inches(0.4))
                card = box(s, x, Inches(1.7), colw, Inches(4.9), c["card"], MSO_SHAPE.ROUNDED_RECTANGLE)
                card.adjustments[0] = 0.06
                box(s, x, Inches(1.7), colw, Inches(0.12), c["accent"] if i == 0 else c["accent2"])
                text(s, x + Inches(0.35), Inches(2.0), colw - Inches(0.7), Inches(0.7), col["titulo"], 24, c["text"],
                     bold=True)
                bullets(s, x + Inches(0.35), Inches(2.8), colw - Inches(0.7), Inches(3.6), col["topicos"], 21)
        else:  # topicos
            header(s, sl["titulo"], n)
            box(s, W - Inches(3.4), Inches(1.7), Inches(2.8), Inches(4.9), c["card"], MSO_SHAPE.ROUNDED_RECTANGLE)
            text(s, W - Inches(3.4), Inches(2.4), Inches(2.8), Inches(1.6), f"{n:02d}", 80, c["accent"], bold=True,
                 align=PP_ALIGN.CENTER)
            if sl["destaque"]:
                text(s, W - Inches(3.2), Inches(4.1), Inches(2.4), Inches(2.2), sl["destaque"], 16, c["muted"],
                     align=PP_ALIGN.CENTER)
            bullets(s, Inches(0.9), Inches(1.8), W - Inches(4.8), Inches(4.8), sl["topicos"],
                    26 if len(sl["topicos"]) <= 3 else 23 if len(sl["topicos"]) <= 4 else 20)
        notes(s, sl["notas"])

    if plan.get("referencias"):
        s = prs.slides.add_slide(blank)
        background(s)
        header(s, "Referências", len(plan["slides"]) + 1)
        bullets(s, Inches(0.9), Inches(1.8), W - Inches(1.8), Inches(4.8), plan["referencias"], 18)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return str(path)


def pptx_to_pdf(path):
    """O PDF do .pptx: pelo PowerPoint (se tiver) ou pelo LibreOffice. None se não der."""
    path = Path(path)
    pdf = path.with_suffix(".pdf")
    try:
        import comtypes.client
        app = comtypes.client.CreateObject("PowerPoint.Application")
        pres = app.Presentations.Open(str(path), WithWindow=False)
        pres.SaveAs(str(pdf), 32)  # 32 = ppSaveAsPDF
        pres.Close()
        if pdf.is_file():
            return str(pdf)
    except Exception:
        pass
    import shutil
    import subprocess
    office = shutil.which("soffice") or next((str(p) for p in (Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
                                                               Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"))
                                              if p.is_file()), "")
    if office:
        try:
            subprocess.run([office, "--headless", "--convert-to", "pdf", "--outdir", str(path.parent), str(path)],
                           capture_output=True, timeout=180, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            if pdf.is_file():
                return str(pdf)
        except Exception:
            pass
    return None


# ========================= GOOGLE SLIDES =========================

SLIDES_API = "https://slides.googleapis.com/v1/presentations"
GOOGLE_LAYOUT = {"secao": "SECTION_HEADER", "destaque": "BIG_NUMBER", "citacao": "MAIN_POINT",
                 "duas_colunas": "TITLE_AND_TWO_COLUMNS", "topicos": "TITLE_AND_BODY"}


def google_requests(plan):
    """Os pedidos (batchUpdate) que montam a apresentação no layout padrão do Google Slides."""
    reqs, n = [], 0

    def slide(layout, mapping):
        nonlocal n
        n += 1
        sid = f"bolha_s{n}"
        ids = {}
        maps = []
        for kind, index in mapping:
            oid = f"{sid}_{kind.lower()}_{index}"
            ids[(kind, index)] = oid
            maps.append({"layoutPlaceholder": {"type": kind, "index": index}, "objectId": oid})
        reqs.append({"createSlide": {"objectId": sid, "slideLayoutReference": {"predefinedLayout": layout},
                                     "placeholderIdMappings": maps}})
        return ids

    def put(oid, value):
        if value:
            reqs.append({"insertText": {"objectId": oid, "text": str(value)}})

    ids = slide("TITLE", [("CENTERED_TITLE", 0), ("SUBTITLE", 0)])
    put(ids[("CENTERED_TITLE", 0)], plan["titulo"])
    put(ids[("SUBTITLE", 0)], plan.get("subtitulo") or datetime.now().strftime("%d/%m/%Y"))
    for sl in plan["slides"]:
        layout = GOOGLE_LAYOUT[sl["layout"]]
        if layout == "TITLE_AND_TWO_COLUMNS":
            ids = slide(layout, [("TITLE", 0), ("BODY", 0), ("BODY", 1)])
            put(ids[("TITLE", 0)], sl["titulo"])
            for i, col in enumerate(sl["colunas"][:2]):
                put(ids[("BODY", i)], "\n".join([col["titulo"]] + col["topicos"]))
        elif layout == "BIG_NUMBER":
            ids = slide(layout, [("TITLE", 0), ("BODY", 0)])
            put(ids[("TITLE", 0)], sl["destaque"])
            put(ids[("BODY", 0)], " · ".join(filter(None, [sl["titulo"], sl["legenda"]])))
        elif layout == "MAIN_POINT":
            ids = slide(layout, [("TITLE", 0)])
            put(ids[("TITLE", 0)], f"“{sl['destaque']}”" + (f"\n— {sl['legenda']}" if sl["legenda"] else ""))
        elif layout == "SECTION_HEADER":
            ids = slide(layout, [("TITLE", 0)])
            put(ids[("TITLE", 0)], sl["titulo"])
        else:
            ids = slide(layout, [("TITLE", 0), ("BODY", 0)])
            put(ids[("TITLE", 0)], sl["titulo"])
            put(ids[("BODY", 0)], "\n".join(sl["topicos"]))
    if plan.get("referencias"):
        ids = slide("TITLE_AND_BODY", [("TITLE", 0), ("BODY", 0)])
        put(ids[("TITLE", 0)], "Referências")
        put(ids[("BODY", 0)], "\n".join(plan["referencias"]))
    return reqs


def make_google(plan, google, download, pdf_path):
    """Cria no Google Slides (layout padrão do Google) e baixa o PDF. Devolve (link, pdf ou None)."""
    created = google(SLIDES_API, {"title": plan["titulo"]})
    pid = created["presentationId"]
    first = [s["objectId"] for s in created.get("slides", [])]
    reqs = google_requests(plan) + [{"deleteObject": {"objectId": oid}} for oid in first]
    google(f"{SLIDES_API}/{pid}:batchUpdate", {"requests": reqs})
    link = f"https://docs.google.com/presentation/d/{pid}/edit"
    pdf = None
    try:
        data = download(f"https://www.googleapis.com/drive/v3/files/{pid}/export?mimeType=application/pdf")
        Path(pdf_path).parent.mkdir(parents=True, exist_ok=True)
        Path(pdf_path).write_bytes(data)
        pdf = str(pdf_path)
    except Exception:
        pass
    return link, pdf


# ========================= O PEDIDO INTEIRO =========================

def make_slides(topic, count, level, style, where, folder, think, snippets=None, google=None,
                download=None, google_ready=False, extra=""):
    """Faz a apresentação. Devolve {"pptx", "pdf", "link", "titulo", "slides"} ou None se não deu."""
    topic = str(topic or "").strip()
    if not topic:
        _say("Sobre qual tema é o slide?")
        return None
    try:
        count = max(3, min(25, int(count or 8)))
    except (TypeError, ValueError):
        count = 8
    _say(f"Montando {count} slides sobre {topic}. Leva uns segundos...")
    try:
        plan = outline(topic, count, level, think, snippets, extra)
    except Exception as e:
        _say(f"Não consegui montar o roteiro: {e}.")
        return None
    name = safe_name(plan["titulo"])
    folder = Path(folder)
    out = {"titulo": plan["titulo"], "slides": len(plan["slides"]) + 1, "pptx": None, "pdf": None, "link": None}
    if norm(where) == "google":
        if not google_ready:
            _say("O Google ainda não está conectado: faço no PowerPoint. Para usar o Google Slides, diga \"conecta o "
                 "Google\".")
        else:
            try:
                out["link"], out["pdf"] = make_google(plan, google, download, folder / f"{name}.pdf")
            except Exception as e:
                _say(f"O Google Slides não deixou criar ({e}); faço no PowerPoint.")
    if not out["link"]:
        try:
            out["pptx"] = make_pptx(plan, folder / f"{name}.pptx", style)
        except ImportError:
            _say("Para montar o PowerPoint instale: python -m pip install python-pptx")
            return None
        out["pdf"] = pptx_to_pdf(out["pptx"])
    where_ = out["link"] or out["pptx"]
    _say(f"Pronto: {out['slides']} slides sobre {plan['titulo']}. Salvei em {where_}"
         + (" e o PDF ao lado." if out["pdf"] and not out["link"] else "."))
    return out
